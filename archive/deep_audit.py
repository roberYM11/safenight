#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.12"
# dependencies = ["numpy", "openpyxl"]
# ///
# ─── How to run ───
# Install uv: https://docs.astral.sh/uv/getting-started/installation/
# Run from the workspace: uv run work/deep_audit.py
# Or: chmod +x work/deep_audit.py && ./work/deep_audit.py
# ──────────────────
"""Reconstruct SafeNight source rows and identify field misalignment."""

from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def main() -> None:
    """Audit source lineage and write diagnostic inputs without altering originals."""
    root = Path(__file__).resolve().parent
    source = root / "sources"
    out = root / "deep-audit"
    out.mkdir(exist_ok=True)
    with (source / "SDot_and_Loc.csv").open(newline="") as handle:
        old = list(csv.DictReader(handle))
    with (source / "test_SDOT_ALL.csv").open(newline="") as handle:
        final = list(csv.DictReader(handle))
    wb = load_workbook(source / "raw_locations.xlsx", read_only=True, data_only=True)
    locations: defaultdict[str, list[tuple[float, float]]] = defaultdict(list)
    for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        if row[1] and row[4] is not None and row[5] is not None:
            locations[str(row[1])].append((float(row[4]), float(row[5])))
    wb.close()
    with (source / "raw_sdot.csv").open(encoding="cp949", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        raw = list(reader)
    widths = Counter(map(len, raw))
    tails = sum(any(value.strip() for value in row[len(header):]) for row in raw)
    selected = []
    merged_index = 0
    for raw_index, row in enumerate(raw):
        matches = locations.get(row[1], [])
        for lat, lon in matches:
            time = row[2]
            if time.startswith("2025-10-12_") and time[11:13] in {"00", "01", "02", "22", "23"}:
                if all(row[i].strip() for i in [24, 25, 26, 30, 31, 32]):
                    selected.append((merged_index, raw_index, row, lat, lon))
            merged_index += 1
        if not matches:
            merged_index += 1
    lineage_errors = 0
    feature_errors = 0
    candidates = []
    for saved, target, item in zip(old, final, selected, strict=True):
        joined, raw_index, row, lat, lon = item
        same_location = abs(float(saved["위도"]) - lat) < 1e-11 and abs(float(saved["경도"]) - lon) < 1e-11
        same_time = saved["level_2"] == row[2].replace("_", " ")
        lineage_errors += int(not (int(saved[""]) == joined and same_time and same_location))
        feature_errors += int(any(abs(float(saved[c]) - float(row[i])) > 1e-12 for c, i in [
            ("조도 최대(lux)", 24), ("조도 평균(lux)", 25), ("조도 최소(lux)", 26),
            ("소음 최대(dB)", 30), ("소음 평균(dB)", 31), ("소음 최소(dB)", 32),
        ]))
        real = [row[i] for i in [18, 19, 20, 24, 25, 26]]
        if all(value.strip() for value in real):
            values = [float(value) for value in real]
            candidates.append([raw_index, row[1], row[2], lat, lon, *values,
                               (values[3] - values[5]) / 4,
                               *[float(target[c]) for c in ["CCTV_개수_25m", "가로등_개수_25m", "보안등_개수_25m"]],
                               float(target["safety_index"])])
    names = ["raw_index", "sensor_id", "time", "lat", "lon", "lux_max", "lux_avg", "lux_min",
             "noise_max", "noise_avg", "noise_min", "noise_range_proxy", "cctv", "light", "security", "old_target"]
    with (out / "corrected_inputs.csv").open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(names)
        writer.writerows(candidates)
    numeric = np.array([row[5:] for row in candidates], dtype=float)
    summary = {
        "raw_rows": len(raw), "header_width": len(header), "row_width_counts": dict(widths),
        "nonempty_extra_fields": tails, "selected_reconstructed": len(selected),
        "lineage_error_rows": lineage_errors, "six_shift_feature_error_rows": feature_errors,
        "location_serial_count": len(locations), "location_duplicate_serials": sum(len(v) > 1 for v in locations.values()),
        "corrected_complete_rows": len(candidates),
        "raw_sensor_time_duplicate_excess": len(selected) - len({(r[2][1], r[2][2]) for r in selected}),
        "corrected_ranges": {name: {"min": float(numeric[:, i].min()), "max": float(numeric[:, i].max()),
                                   "median": float(np.median(numeric[:, i])), "unique": int(len(np.unique(numeric[:, i])))}
                             for i, name in enumerate(names[5:11])},
        "field_mapping": {"stored_lux": "raw_noise", "stored_noise_before_x1000": "raw_vibration_y"},
    }
    (out / "raw-audit.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
